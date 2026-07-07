from __future__ import annotations

import io

import pytest
from pptx import Presentation
from pptx.util import Inches

from ingestor.slides import extract_slides

pytestmark = pytest.mark.seam1


def _deck_bytes(png_bytes: bytes) -> bytes:
    presentation = Presentation()
    layout = presentation.slide_layouts[5]

    slide_one = presentation.slides.add_slide(layout)
    slide_one.shapes.title.text = "Q3 Roadmap"
    text_box = slide_one.shapes.add_textbox(Inches(1), Inches(2), Inches(4), Inches(1))
    text_box.text_frame.text = "Ship v2 by August"

    slide_two = presentation.slides.add_slide(layout)
    slide_two.shapes.title.text = "Appendix"
    slide_two.shapes.add_picture(io.BytesIO(png_bytes), Inches(1), Inches(1), height=Inches(2))

    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def _tiny_png() -> bytes:
    from PIL import Image

    buffer = io.BytesIO()
    Image.new("RGB", (2, 2), color="white").save(buffer, format="PNG")
    return buffer.getvalue()


def test_splits_a_deck_into_slides_in_order() -> None:
    slides = extract_slides(_deck_bytes(_tiny_png()))

    assert len(slides) == 2
    assert "Q3 Roadmap" in slides[0].text
    assert "Ship v2 by August" in slides[0].text
    assert slides[0].images == []


def test_extracts_embedded_pictures_from_a_slide() -> None:
    slides = extract_slides(_deck_bytes(_tiny_png()))

    assert len(slides[1].images) == 1
    assert slides[1].images[0].extension == "png"
    assert slides[1].images[0].data
