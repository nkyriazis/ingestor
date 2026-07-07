from __future__ import annotations

import io
from dataclasses import dataclass, field

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


@dataclass
class ExtractedImage:
    data: bytes
    extension: str


@dataclass
class ExtractedSlide:
    text: str
    images: list[ExtractedImage] = field(default_factory=list)


def extract_slides(raw_bytes: bytes) -> list[ExtractedSlide]:
    """Splits a slide deck (.pptx) into its individual slides, in order —
    the structure discovery CONTEXT.md's Conversion entry describes. Slide
    text is read natively (accurate, no OCR needed); embedded pictures are
    returned as raw bytes for the caller to caption separately, since each
    is its own Evidence node.
    """
    presentation = Presentation(io.BytesIO(raw_bytes))
    slides = []
    for slide in presentation.slides:
        texts = [
            shape.text_frame.text
            for shape in slide.shapes
            if shape.has_text_frame and shape.text_frame.text.strip()
        ]
        images = [
            ExtractedImage(data=shape.image.blob, extension=shape.image.ext)
            for shape in slide.shapes
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
        ]
        slides.append(ExtractedSlide(text="\n".join(texts), images=images))
    return slides
